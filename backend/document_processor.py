"""Document processing: extract text AND images from PDF, DOCX, TXT, and images."""
import os
import io
import fitz  # PyMuPDF
from docx import Document as DocxDocument
from PIL import Image
import pytesseract
from pathlib import Path
from typing import Optional
from loguru import logger

from config import settings

if settings.TESSERACT_CMD and os.path.exists(settings.TESSERACT_CMD):
    pytesseract.pytesseract.tesseract_cmd = settings.TESSERACT_CMD


class DocumentProcessor:
    """Extract text content and images from various file formats."""

    SUPPORTED_TYPES = {
        ".pdf": "pdf",
        ".docx": "docx",
        ".doc": "docx",
        ".txt": "text",
        ".md": "text",
        ".png": "image",
        ".jpg": "image",
        ".jpeg": "image",
        ".tiff": "image",
        ".bmp": "image",
        ".webp": "image",
    }

    def __init__(self):
        self.chunk_size = settings.CHUNK_SIZE
        self.chunk_overlap = settings.CHUNK_OVERLAP

    def get_file_type(self, filename: str) -> Optional[str]:
        ext = Path(filename).suffix.lower()
        return self.SUPPORTED_TYPES.get(ext)

    def extract_text(self, file_path: str, file_type: str) -> str:
        """Extract text from file based on type."""
        try:
            if file_type == "pdf":
                return self._extract_pdf(file_path)
            elif file_type == "docx":
                return self._extract_docx(file_path)
            elif file_type == "text":
                return self._extract_text(file_path)
            elif file_type == "image":
                return self._extract_image(file_path)
            else:
                raise ValueError(f"Unsupported file type: {file_type}")
        except Exception as e:
            logger.error(f"Error extracting text from {file_path}: {e}")
            raise

    def extract_images(self, file_path: str, file_type: str) -> list[dict]:
        """Extract images from file. Returns list of {image_path, page_num, width, height}."""
        images = []
        try:
            if file_type == "pdf":
                images = self._extract_pdf_images(file_path)
            elif file_type == "image":
                img = Image.open(file_path)
                images.append({
                    "image_path": file_path,
                    "page_num": 0,
                    "width": img.width,
                    "height": img.height,
                })
            elif file_type == "docx":
                images = self._extract_docx_images(file_path)
        except Exception as e:
            logger.error(f"Error extracting images from {file_path}: {e}")
        return images

    def _extract_pdf(self, file_path: str) -> str:
        """Extract text from PDF using PyMuPDF."""
        text_parts = []
        with fitz.open(file_path) as doc:
            for page_num, page in enumerate(doc):
                text = page.get_text()
                if text.strip():
                    text_parts.append(f"[Page {page_num + 1}]\n{text}")
        return "\n\n".join(text_parts)

    def _extract_pdf_images(self, file_path: str) -> list[dict]:
        """Extract embedded images from PDF."""
        images = []
        with fitz.open(file_path) as doc:
            for page_num, page in enumerate(doc):
                image_list = page.get_images(full=True)
                for img_index, img in enumerate(image_list):
                    xref = img[0]
                    base_image = doc.extract_image(xref)
                    image_bytes = base_image["image"]
                    image_ext = base_image["ext"]

                    # Save extracted image
                    img_filename = f"pdf_p{page_num+1}_img{img_index+1}.{image_ext}"
                    img_path = settings.UPLOAD_DIR / "extracted_images" / img_filename
                    img_path.parent.mkdir(parents=True, exist_ok=True)

                    with open(img_path, "wb") as f:
                        f.write(image_bytes)

                    # Get image dimensions
                    try:
                        pil_img = Image.open(io.BytesIO(image_bytes))
                        width, height = pil_img.width, pil_img.height
                    except Exception:
                        width, height = 0, 0

                    images.append({
                        "image_path": str(img_path),
                        "page_num": page_num + 1,
                        "width": width,
                        "height": height,
                    })
        return images

    def _extract_docx(self, file_path: str) -> str:
        """Extract text from DOCX file."""
        doc = DocxDocument(file_path)
        text_parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text for cell in row.cells)
                if row_text.strip():
                    text_parts.append(row_text)
        return "\n".join(text_parts)

    def _extract_docx_images(self, file_path: str) -> list[dict]:
        """Extract images from DOCX file."""
        images = []
        doc = DocxDocument(file_file)
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                image_part = rel.target_part
                image_bytes = image_part.blob
                image_ext = Path(rel.target_ref).suffix

                img_filename = f"docx_img_{len(images)+1}{image_ext}"
                img_path = settings.UPLOAD_DIR / "extracted_images" / img_filename
                img_path.parent.mkdir(parents=True, exist_ok=True)

                with open(img_path, "wb") as f:
                    f.write(image_bytes)

                try:
                    pil_img = Image.open(io.BytesIO(image_bytes))
                    width, height = pil_img.width, pil_img.height
                except Exception:
                    width, height = 0, 0

                images.append({
                    "image_path": str(img_path),
                    "page_num": 0,
                    "width": width,
                    "height": height,
                })
        return images

    def _extract_text(self, file_path: str) -> str:
        """Extract text from plain text file."""
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    def _extract_image(self, file_path: str) -> str:
        """Extract text from image using OCR."""
        image = Image.open(file_path)
        if image.mode not in ("RGB", "L"):
            image = image.convert("RGB")
        text = pytesseract.image_to_string(image, lang=settings.OCR_LANG)
        return text.strip()

    def chunk_text(self, text: str) -> list[str]:
        """Split text into overlapping chunks."""
        if not text.strip():
            return []

        chunks = []
        start = 0
        text_len = len(text)

        while start < text_len:
            end = start + self.chunk_size
            if end < text_len:
                boundary = text.rfind(". ", start, end)
                if boundary == -1 or boundary < start + self.chunk_size // 2:
                    boundary = text.rfind(" ", start, end)
                if boundary > start:
                    end = boundary + 1

            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)

            start = end - self.chunk_overlap
            if start < 0:
                start = 0

        return chunks

    def process_file(self, file_path: str, filename: str) -> dict:
        """Process a file and return extracted content with metadata."""
        file_type = self.get_file_type(filename)
        if not file_type:
            raise ValueError(f"Unsupported file type: {filename}")

        text = self.extract_text(file_path, file_type)
        chunks = self.chunk_text(text)
        images = self.extract_images(file_path, file_type)

        file_size = os.path.getsize(file_path)

        return {
            "filename": filename,
            "file_type": file_type,
            "file_path": file_path,
            "file_size": file_size,
            "content_text": text,
            "chunks": chunks,
            "chunk_count": len(chunks),
            "images": images,
            "image_count": len(images),
            "metadata": {
                "file_type": file_type,
                "file_size": file_size,
                "chunk_count": len(chunks),
                "image_count": len(images),
            },
        }
